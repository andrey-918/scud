import pandas as pd
from tkinter import messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Border, Side, Alignment, Font
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.formatting.rule import CellIsRule
from config.constants import DATABASE_VISITS, DATABASE_VISITS_REPORT, DATABASE_REQUESTS, REPORTS_FOLDER, SHORT_NAMES, DEFAULT_MIN_PERCENT
from config.settings import load_settings
from database.db_init import create_connection
from utils.file_utils import delete_old_reports
from utils.time_utils import save_last_report_date
from datetime import datetime

def generate_visits_report(report_date=None, root=None):
    try:
        if report_date is None:
            report_date = datetime.now().date()

        conn_visits = create_connection(DATABASE_VISITS)
        if conn_visits:
            visits_df = pd.read_sql_query("SELECT * FROM visits", conn_visits)
            visits_df.rename(columns=SHORT_NAMES, inplace=True)

            date_str = report_date.strftime("%Y-%m-%d")
            default_filename = f"Отчет_по_посещениям_{date_str}.xlsx"
            report_path = os.path.join(REPORTS_FOLDER, default_filename)

            wb = Workbook()
            ws = wb.active

            for r in dataframe_to_rows(visits_df, index=False, header=True):
                ws.append(r)

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    if cell.row == 1 or cell.column_letter in ["A", "B", "C"]:
                        cell.border = thin_border

            for col in ws.columns:
                max_length = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max_length + 2

            wb.save(report_path)
            if root:
                messagebox.showinfo("Успех", f"Отчёт №1 сохранён в {report_path}", parent=root)

            with open("last_report_date.txt", "w") as file:
                file.write(date_str)

            cursor = conn_visits.cursor()
            columns_to_reset = [
                "Понедельник_Завтрак", "Понедельник_Обед", "Понедельник_Ужин",
                "Вторник_Завтрак", "Вторник_Обед", "Вторник_Ужин",
                "Среда_Завтрак", "Среда_Обед", "Среда_Ужин",
                "Четверг_Завтрак", "Четверг_Обед", "Четверг_Ужин",
                "Пятница_Завтрак", "Пятница_Обед", "Пятница_Ужин",
                "Суббота_Завтрак", "Суббота_Обед", "Суббота_Ужин"
            ]
            reset_query = f"UPDATE visits SET {', '.join([f'{col} = 0' for col in columns_to_reset])}"
            cursor.execute(reset_query)
            conn_visits.commit()

            save_last_report_date(report_date)
            delete_old_reports(REPORTS_FOLDER)
    except Exception as e:
        if root:
            messagebox.showerror("Ошибка", f"Не удалось сформировать отчёт №1: {e}", parent=root)
        else:
            print(f"Ошибка при формировании отчёта №1: {e}")
    finally:
        if conn_visits:
            conn_visits.close()

def generate_visits_report_everyday(root):
    try:
        conn_visits = create_connection(DATABASE_VISITS)
        if conn_visits:
            visits_df = pd.read_sql_query("SELECT * FROM visits", conn_visits)
            visits_df.rename(columns=SHORT_NAMES, inplace=True)

            current_date = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            default_filename = f"Отчет_по_посещениям_РГ_{current_date}.xlsx"
            report_path = os.path.join(REPORTS_FOLDER, default_filename)

            wb = Workbook()
            ws = wb.active

            for r in dataframe_to_rows(visits_df, index=False, header=True):
                ws.append(r)

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                for cell in row:
                    if cell.row == 1 or cell.column_letter in ["A", "B", "C"]:
                        cell.border = thin_border

            for col in ws.columns:
                max_length = max(len(str(cell.value or "")) for cell in col)
                ws.column_dimensions[col[0].column_letter].width = max_length + 2

            wb.save(report_path)
            messagebox.showinfo("Успех", f"Отчёт №1 сохранён в {report_path}", parent=root)

            delete_old_reports(REPORTS_FOLDER)
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось сформировать отчёт №1: {e}", parent=root)
    finally:
        if conn_visits:
            conn_visits.close()

def generate_analytics_report(root):
    try:
        conn_visits = create_connection(DATABASE_VISITS_REPORT)
        conn_requests = create_connection(DATABASE_REQUESTS)

        if conn_visits and conn_requests:
            visits_df = pd.read_sql_query("SELECT * FROM visits_report", conn_visits)
            requests_df = pd.read_sql_query("SELECT * FROM requests", conn_requests)

            report_df = pd.merge(
                visits_df,
                requests_df,
                on=["uid", "name", "student_group"],
                suffixes=("_visit", "_request")
            )

            meal_columns = [
                "Понедельник_Завтрак", "Понедельник_Обед", "Понедельник_Ужин",
                "Вторник_Завтрак", "Вторник_Обед", "Вторник_Ужин",
                "Среда_Завтрак", "Среда_Обед", "Среда_Ужин",
                "Четверг_Завтрак", "Четверг_Обед", "Четверг_Ужин",
                "Пятница_Завтрак", "Пятница_Обед", "Пятница_Ужин",
                "Суббота_Завтрак", "Суббота_Обед", "Суббота_Ужин"
            ]

            rename_dict = {col + "_visit": col for col in meal_columns}
            report_df.rename(columns=rename_dict, inplace=True)
            report_df.drop(columns=[col + "_request" for col in meal_columns], inplace=True)

            report_df["Всего заявок"] = requests_df[meal_columns].sum(axis=1)
            report_df["Посещения по заявке"] = report_df.apply(
                lambda row: sum(1 for col in meal_columns if requests_df.loc[row.name, col] == 1 and row[col] == 1),
                axis=1
            )
            report_df["Посещения без заявки"] = report_df.apply(
                lambda row: sum(1 for col in meal_columns if requests_df.loc[row.name, col] == 0 and row[col] == 1),
                axis=1
            )
            report_df["Процент посещения"] = report_df.apply(
                lambda row: round((row["Посещения по заявке"] / row["Всего заявок"] * 100))
                if row["Всего заявок"] > 0 else 0,
                axis=1
            )

            final_columns = ["name", "student_group", *meal_columns, "Всего заявок", "Посещения по заявке",
                             "Посещения без заявки", "Процент посещения"]
            final_report_df = report_df[final_columns].copy()
            final_report_df.rename(columns={"name": "ФИО", "student_group": "Группа"}, inplace=True)
            final_report_df.rename(columns=SHORT_NAMES, inplace=True)

            total_row = {"ФИО": "Итого", "Группа": ""}
            for col in meal_columns:
                renamed_col = SHORT_NAMES.get(col, col)
                total_row[renamed_col] = final_report_df[renamed_col].sum()
            total_row["Всего заявок"] = ""
            total_row["Посещения по заявке"] = ""
            total_row["Посещения без заявки"] = ""
            total_row["Процент посещения"] = ""
            total_df = pd.DataFrame([total_row])
            final_report_df = pd.concat([final_report_df, total_df], ignore_index=True)

            settings = load_settings()
            min_percent = settings.get("min_percent", DEFAULT_MIN_PERCENT)

            current_date = datetime.now().strftime("%Y-%m-%d")
            default_filename = f"Отчет_с_аналитикой_{current_date}.xlsx"
            report_path = filedialog.asksaveasfilename(
                defaultextension=".xlsx",
                filetypes=[("Excel files", "*.xlsx")],
                initialfile=default_filename
            )
            if report_path:
                wb = Workbook()
                ws = wb.active

                for r in dataframe_to_rows(final_report_df, index=False, header=True):
                    ws.append(r)

                yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
                thin_border = Border(
                    left=Side(style='thin'),
                    right=Side(style='thin'),
                    top=Side(style='thin'),
                    bottom=Side(style='thin')
                )

                for row in ws.iter_rows(min_row=2, max_row=ws.max_row - 1, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        if cell.column_letter in ["C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M", "N", "O", "P", "Q", "R", "S"]:
                            if cell.value == 1 and requests_df.loc[cell.row - 2, meal_columns[cell.column - 3]] == 0:
                                cell.fill = yellow_fill

                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        if cell.row == 1 or cell.column_letter in ["A", "B", "U", "V", "W", "X"] or cell.row == ws.max_row:
                            cell.border = thin_border

                red_font = Font(color="FF0000")
                conditional_rule = CellIsRule(operator='lessThan', formula=[min_percent], font=red_font)
                ws.conditional_formatting.add(f"X2:X{ws.max_row - 1}", conditional_rule)

                for col in ws.columns:
                    max_length = max(len(str(cell.value or "")) for cell in col)
                    ws.column_dimensions[col[0].column_letter].width = max_length + 2

                ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=2)
                ws.merge_cells(start_row=ws.max_row, start_column=21, end_row=ws.max_row, end_column=24)
                ws.cell(row=ws.max_row, column=1).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=ws.max_row, column=21).alignment = Alignment(horizontal="center", vertical="center")

                for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
                    for cell in row:
                        cell.alignment = Alignment(horizontal="center", vertical="center")

                wb.save(report_path)
                messagebox.showinfo("Успех", f"Отчёт №2 сохранён в {report_path}", parent=root)
                delete_old_reports(REPORTS_FOLDER)
    except Exception as e:
        messagebox.showerror("Ошибка", f"Не удалось сформировать отчёт №2: {e}", parent=root)
    finally:
        if conn_visits:
            conn_visits.close()
        if conn_requests:
            conn_requests.close()